# -*- coding: utf-8 -*-
"""
Extração FIDC via API do FNET — Interface local (Netz Asset)
============================================================

App web local: inicia um pequeno servidor no seu computador e abre uma página
no navegador. Você preenche o formulário (nome OU CNPJ do fundo, caminho do
Excel, últimos meses, sobrescrever, backup) e acompanha o progresso ao vivo.

Como usar
---------
    python fnet_app.py

O navegador abre sozinho em http://127.0.0.1:<porta>. Se não abrir, copie o
endereço mostrado no terminal.

Requisitos: Python 3.8+ e a biblioteca openpyxl (pip install openpyxl).
Nada de Playwright/Chromium. O tkinter (para o botão "Procurar...") já vem com
o Python padrão no Windows.

Esta interface reaproveita exatamente a mesma lógica do notebook
"extracao_fnet_fidc_api (OFICIAL).ipynb".
"""

import json, time, re, os, sys, zipfile, shutil, threading, webbrowser
import urllib.parse, urllib.request, urllib.error, http.cookiejar
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERRO: a biblioteca 'openpyxl' não está instalada.")
    print("Instale com:  pip install openpyxl")
    sys.exit(1)


# ============================================================================
#  1. Cliente da API do FNET (urllib, com retry para o Cloudflare)
# ============================================================================
BASE = "https://fnet.bmfbovespa.com.br/fnet/publico/"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/120 Safari/537.36")


class FnetClient:
    def __init__(self, pausa=0.5):
        self.pausa = pausa
        self.cj = http.cookiejar.CookieJar()
        self.op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.cj))
        self.op.addheaders = [
            ("User-Agent", UA),
            ("Accept", "application/json, text/javascript, */*; q=0.01"),
            ("Accept-Language", "pt-BR,pt;q=0.9"),
            ("X-Requested-With", "XMLHttpRequest"),
            ("Referer", BASE + "abrirGerenciadorDocumentosCVM"),
        ]

    def _get(self, path, params=None, tentativas=6, espera=3.0):
        url = BASE + path
        if params:
            url += "?" + urllib.parse.urlencode(params, doseq=True)
        time.sleep(self.pausa)
        ultimo = ""
        for t in range(tentativas):
            try:
                with self.op.open(url, timeout=60) as r:
                    if r.status == 200:
                        return r.read()
                    ultimo = f"HTTP {r.status}"
            except urllib.error.HTTPError as e:
                ultimo = f"HTTP {e.code}"
            except Exception as e:
                ultimo = str(e)
            time.sleep(espera * (t + 1))
        raise RuntimeError(f"falha em {path} após {tentativas} tentativas ({ultimo})")

    def iniciar_sessao(self):
        self._get("abrirGerenciadorDocumentosCVM")

    def _listar_fundos(self, termo):
        raw = self._get("listarFundos", {"term": termo, "page": 1,
                                         "idTipoFundo": 2, "idAdm": 0, "paraCerts": "false"})
        return json.loads(raw).get("results", [])

    def resolver_fundo(self, termo, modo="nome"):
        """Resolve o fundo para (idFundo, nome_exato).
        modo="nome": busca pelo nome (pode ser ambíguo -> exige nome exato entre os candidatos).
        modo="cnpj": busca pelo CNPJ (o FNET aceita o CNPJ como termo, mas SÓ com dígitos —
                     com pontuação retorna vazio; por isso removemos qualquer não-dígito)."""
        modo = (modo or "nome").strip().lower()
        if modo == "cnpj":
            digitos = re.sub(r"\D", "", termo)
            if len(digitos) != 14:
                raise RuntimeError(
                    f"CNPJ inválido: '{termo}'. Informe os 14 dígitos "
                    "(ex.: 23.216.398/0001-01 ou 23216398000101).")
            res = self._listar_fundos(digitos)
            if not res:
                raise RuntimeError(f"nenhum fundo encontrado para o CNPJ '{termo}'.")
            if len(res) == 1:
                return res[0]["id"], res[0]["text"]
            opcoes = "\n".join(f"  - {x['text']}" for x in res)
            raise RuntimeError(f"CNPJ ambíguo '{termo}'. Candidatos:\n{opcoes}")

        nome = termo
        res = self._listar_fundos(nome)
        if not res:
            raise RuntimeError(f"nenhum fundo encontrado para '{nome}'.")
        exatos = [x for x in res if x["text"].strip().upper() == nome.strip().upper()]
        if len(exatos) == 1:
            return exatos[0]["id"], exatos[0]["text"]
        if len(res) == 1:
            return res[0]["id"], res[0]["text"]
        opcoes = "\n".join(f"  - {x['text']}" for x in res)
        raise RuntimeError(f"nome ambíguo '{nome}'. Candidatos:\n{opcoes}\n"
                           f"Use o nome EXATO de um deles, ou busque pelo CNPJ.")

    def listar_informes(self, id_fundo, nome_exato=None):
        """Retorna {(ano, mes): id_documento} dos Informes Mensais (maior versão por mês).
        Filtra pelo idFundo (o servidor já restringe ao fundo) e pelo tipo de documento.
        NÃO exige que o nome do documento seja idêntico ao resolvido: em fundos com
        estrutura de classe o nome no listarFundos (ex.: 'FDC UNIVERSI - PRAVALER…')
        difere do descricaoFundo dos informes (ex.: 'PRAVALER…'), e exigir igualdade
        descartaria informes válidos."""
        PAGINA = 200
        linhas, inicio = [], 0
        while True:
            params = {"d": 1, "s": inicio, "l": PAGINA, "q": "", "o[0][dataEntrega]": "desc",
                      "tipoFundo": 2, "idFundo": id_fundo, "situacao": "A"}
            j = json.loads(self._get("pesquisarGerenciadorDocumentosDados", params))
            data = j.get("data", [])
            linhas.extend(data)
            inicio += PAGINA
            if not data or inicio >= int(j.get("recordsFiltered", 0)):
                break
        inf = [x for x in linhas
               if "informe mensal" in (x.get("tipoDocumento") or "").lower()]
        if not inf:
            raise RuntimeError(
                f"nenhum Informe Mensal retornado para idFundo={id_fundo}. "
                "Confira o fundo, ou o servidor do FNET pode estar instável — tente de novo.")
        melhor = {}
        for x in inf:
            ref = x.get("dataReferencia")
            if not ref or "/" not in ref:
                continue
            mes, ano = ref.split("/")[0], ref.split("/")[-1]
            chave = (int(ano), int(mes))
            v = int(x.get("versao") or 0)
            if chave not in melhor or v >= melhor[chave][0]:
                melhor[chave] = (v, x["id"])
        return {k: vid for k, (v, vid) in melhor.items()}

    def baixar_xml(self, id_doc):
        return self._get("downloadDocumento", {"id": id_doc})


# ============================================================================
#  2. Leitura do XML do Informe Mensal FIDC
# ============================================================================
def numero_br(s):
    if not s:
        return 0.0
    s = s.strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _t(node, tag):
    el = node.find(tag) if node is not None else None
    return el.text if (el is not None and el.text) else ""


def _f(node, tag):
    return numero_br(_t(node, tag))


_BUCKETS_90 = ["VL_INAD_VENC_91_120", "VL_INAD_VENC_121_150", "VL_INAD_VENC_151_180",
               "VL_INAD_VENC_181_360", "VL_INAD_VENC_361_720", "VL_INAD_VENC_721_1080",
               "VL_INAD_VENC_1080"]


def _sem_acento(s):
    for de, para in (("á", "a"), ("â", "a"), ("ã", "a"), ("à", "a"),
                     ("é", "e"), ("ê", "e"), ("í", "i"), ("ó", "o"),
                     ("ô", "o"), ("õ", "o"), ("ú", "u"), ("ç", "c")):
        s = s.replace(de, para)
    return s


def classificar_classe(no, tipo, serie):
    """Classifica o rótulo de UMA classe em: MEZANINO, SUBORDINADA, UNICA, SENIOR, OUTRO.

    Os administradores usaram nomes muito diferentes ao longo do tempo — e a
    Resolução CVM 175 reorganizou classes/subclasses em 2024, quebrando a
    nomenclatura de praticamente todos os fundos. Exemplos reais já vistos:
      "Cota Subordinada", "Subordinada 1", "Cota Junior", "Cota Unica",
      "Subordonada" (com erro de digitação), "Única", "Sub | Subordinada 2".
    Por isso o rótulo é apenas um INDÍCIO: a decisão final de qual classe é a
    subordinada é tomada em escolher_serie_coerente(), pela continuidade da
    série ao longo dos meses."""
    txt = _sem_acento(f"{tipo or ''} {serie or ''}".strip().lower())
    if "mezan" in txt or "mezz" in txt:
        return "MEZANINO"
    if no == "SENIOR":
        # o nó XML já diz "sênior"; só um rótulo explícito de subordinada o contradiz
        return "SUBORDINADA" if ("subord" in txt or "junior" in txt) else "SENIOR"
    if "subord" in txt or "junior" in txt or (tipo or "").strip().lower() == "sub":
        return "SUBORDINADA"
    if "unica" in txt or "unico" in txt:
        return "UNICA"
    return "OUTRO"


def _eh_subordinada(tipo, serie):
    """Mantida por compatibilidade: reconhece pelo rótulo, num mês isolado."""
    return classificar_classe("SUBORD", tipo, serie) == "SUBORDINADA"


def escolher_subordinada(classes):
    """Regra só por rótulo, para um mês isolado (usada por parse_informe).
    A seleção boa é a de escolher_serie_coerente(), que olha a série inteira."""
    subs = [c for c in classes if _eh_subordinada(c["TIPO"], c["SERIE"])]
    for c in subs:
        if c["valor"] != 0:
            return c["TIPO"], c["SERIE"]
    if subs:
        return subs[0]["TIPO"], subs[0]["SERIE"]
    nao_mez = [c for c in classes
               if classificar_classe("SUBORD", c["TIPO"], c["SERIE"]) != "MEZANINO"]
    if len(nao_mez) == 1:
        return nao_mez[0]["TIPO"], nao_mez[0]["SERIE"]
    return (None, None)


def extrair_mes(xml_bytes):
    """Lê um Informe Mensal e devolve TODAS as classes declaradas (dos nós
    SENIOR e SUBORD), com quantidade, valor da cota, aporte e resgate de cada
    uma, mais os campos que não dependem da classe escolhida (PL, inadimplência…).

    Não decide qual classe é a subordinada — isso é papel de
    escolher_serie_coerente(), que compara os meses entre si."""
    root = ET.fromstring(xml_bytes)
    d = {"competencia": _t(root.find("CAB_INFORM"), "DT_COMPT")}

    d["pl"] = _f(root.find(".//PATRLIQ"), "VL_PATRIM_LIQ")
    aquis = root.find(".//COMPMT_DICRED_AQUIS")
    sem = root.find(".//COMPMT_DICRED_SEM_AQUIS")
    d["inadimplentes"] = _f(aquis, "VL_SOM_INAD_VENC") + _f(sem, "VL_SOM_INAD_VENC")
    d["inad_over90"] = sum(_f(aquis, b) for b in _BUCKETS_90) + sum(_f(sem, b) for b in _BUCKETS_90)
    d["recompras"] = _f(root.find(".//NEGOC_DICRED_MES/DICRED_MES_ALIEN_RECOMP"), "VL_DICRED_ALIEN")

    def filhos(bloco):
        no = root.find(f".//CAPTA_RESGA_AMORTI/{bloco}")
        return list(no) if no is not None else []

    capt, resg = filhos("CAPT_MES"), filhos("RESG_MES")

    def fluxo(nodes, tag, tipo, serie, n_cands):
        """Captação/resgate do mês para a classe (tipo, serie).
        Devolve (valor, agregado). agregado=True quando o informe traz um único
        nó para todas as subclasses e não há como separar por classe — aí o
        valor é o TOTAL do nó, não o da classe."""
        iguais = [c for c in nodes if c.tag == tag]
        for c in iguais:                       # 1) match exato por (TIPO, SERIE)
            if _t(c, "TIPO") == tipo and _t(c, "SERIE") == serie:
                return _f(c, "VL_TOTAL"), False
        for c in iguais:                       # 2) nó agregado, sem TIPO/SERIE
            if not _t(c, "TIPO") and not _t(c, "SERIE"):
                return _f(c, "VL_TOTAL"), n_cands > 1
        if len(iguais) == 1 and n_cands == 1:   # 3) nó único e classe única: sem ambiguidade
            return _f(iguais[0], "VL_TOTAL"), False
        # 4) o informe atribuiu o fluxo a OUTRA classe (por TIPO/SERIE): esta não teve fluxo.
        #    Importante não cair aqui em "usa o único nó": em MONARCA dez/23 o
        #    RESG_MES traz só 'Sub/Subordinada 1' (o mezanino) com R$ 2,7 mi, e
        #    atribuir isso também à subordinada real distorce tudo.
        return 0.0, False

    dsc = root.find(".//OUTRAS_INFORM/DESC_SERIE_CLASSE")
    cands = []
    for filho in (list(dsc) if dsc is not None else []):
        no = filho.tag.replace("DESC_SERIE_CLASSE_", "")
        if no not in ("SENIOR", "SUBORD"):
            continue
        tipo, serie = _t(filho, "TIPO"), _t(filho, "SERIE")
        cands.append({
            "no": no, "TIPO": tipo, "SERIE": serie,
            "qtd": _f(filho, "QT_COTAS"), "valor": _f(filho, "VL_COTAS"),
            "classe": classificar_classe(no, tipo, serie),
        })

    n_no = {"SENIOR": sum(1 for c in cands if c["no"] == "SENIOR"),
            "SUBORD": sum(1 for c in cands if c["no"] == "SUBORD")}
    for c in cands:
        tag = "CLASSE_SENIOR" if c["no"] == "SENIOR" else "CLASSE_SUBORD"
        c["aporte"], ap_ag = fluxo(capt, tag, c["TIPO"], c["SERIE"], n_no[c["no"]])
        c["resgate"], rg_ag = fluxo(resg, tag, c["TIPO"], c["SERIE"], n_no[c["no"]])
        c["fluxo_agregado"] = ap_ag or rg_ag

    d["candidatos"] = cands
    return d


# ============================================================================
#  2.1  Identificação COERENTE da subordinada ao longo dos meses
# ============================================================================
# O rótulo da classe no informe não é confiável (ver classificar_classe).
# Já a identidade econômica da classe é: ela persiste mês a mês. Usamos três
# invariantes independentes para reconhecer "a mesma classe" em dois meses:
#
#   • valor da cota  — muda só pela rentabilidade do mês. Aporte e resgate NÃO
#                      mexem no valor unitário, só na quantidade.
#   • quantidade     — muda só por aporte/resgate: qtd_nova ≈ qtd + fluxo/valor.
#   • valor total    — quantidade × valor da cota, ajustado pelo fluxo do mês.
#
# Basta UM desses casar bem para a identidade estar estabelecida (por isso o
# erro combinado pesa sobretudo o melhor dos três): quando um administrador
# reagrupa cotas (desdobramento), quantidade e valor unitário mudam juntos mas
# o valor total continua — e vice-versa.
#
# A escolha final é o caminho de MENOR CUSTO ao longo de todos os meses
# (Viterbi), somando o custo de rótulo de cada classe com o custo de
# descontinuidade entre meses consecutivos. Isso é justamente o que se faz a
# olho na planilha: conferir se quantidade e valor da cota fazem sentido com os
# aportes e resgates, e desconfiar de salto brusco sem justificativa.
#
# Calibração dos custos — a propriedade que se quer garantir:
#   Um mês ISOLADO com rótulo trocado pode ser corrigido pela continuidade
#   (corrigi-lo economiza a descontinuidade na entrada E na saída, ~2x
#   PESO_CONTINUIDADE), mas uma sequência LONGA não pode derivar para a classe
#   errada (cada mês a mais custa o rótulo inteiro e economiza no máximo
#   ~1x PESO_CONTINUIDADE). Daí PESO_CONTINUIDADE < CUSTO(MEZANINO) < 2x.
#   Caso real que exige isso: HB CAPITAL set/2025, em que o administrador
#   rotacionou as três classes e a subordinada real saiu como "Mezanino 1".
CUSTO_ROTULO = {"SUBORDINADA": 0.0, "UNICA": 0.15, "OUTRO": 0.6,
                "SENIOR": 20.0, "MEZANINO": 6.0}
CUSTO_ZERADA = 5.0            # classe zerada: perde de um mezanino, mas ganha se não houver mais nada
CUSTO_SENIOR_MONO = 0.3       # fundo mono-classe: o nó SENIOR/SUBORD não é confiável
CUSTO_SENIOR_SO_ELE = 1.5     # no mês, o nó SUBORD está todo zerado e o SENIOR não
PESO_CONTINUIDADE = 4.0       # peso da descontinuidade frente ao custo de rótulo
ERRO_MAX = 3.0                # satura o erro: um mês esquisito não domina o caminho
ERRO_ALERTA = 0.60            # acima disso, avisa "salto sem justificativa"


def _erro_continuidade(a, b, gap=1):
    """Erro normalizado (0 = continuação perfeita) entre a classe escolhida no
    mês anterior (a) e um candidato do mês atual (b). Ver comentário acima."""
    eps = 1e-9
    qa, va, qb, vb = a["qtd"], a["valor"], b["qtd"], b["valor"]

    zer_a = not qa and not va
    zer_b = not qb and not vb
    if zer_a and zer_b:
        return 0.0
    if zer_a or zer_b:
        return 0.5            # classe começou ou encerrou: nem confirma nem desmente

    err_v = abs(vb - va) / max(abs(va), abs(vb), eps)

    ta, tb = qa * va, qb * vb
    v_ref = ((va + vb) / 2.0) or eps

    # O fluxo do mês não é sempre atribuível à classe: há informes com um único
    # nó de captação/resgate para todas as subclasses, e outros que declaram o
    # fluxo só de uma delas. Por isso medimos os dois cenários — com e sem o
    # fluxo — e ficamos com o mais favorável: a variação está justificada SE os
    # aportes/resgates a explicam OU se simplesmente não houve variação. Um
    # salto que nenhuma das duas leituras explica continua caro, que é
    # exatamente o critério usado na conferência manual.
    err_q = err_t = None
    for f in (b["aporte"] - b["resgate"], 0.0):
        eq = abs(qb - (qa + f / v_ref)) / max(abs(qa), abs(qb), eps)
        et = abs(tb - (ta + f)) / max(abs(ta), abs(tb), eps)
        err_q = eq if err_q is None else min(err_q, eq)
        err_t = et if err_t is None else min(err_t, et)

    e = sorted((err_v, err_q, err_t))
    return (0.65 * e[0] + 0.35 * e[1]) / (max(gap, 1) ** 0.5)


def _custo_rotulo(cand, subord_zerado, mono_classe):
    c = CUSTO_ROTULO.get(cand["classe"], 0.6)
    if cand["classe"] == "SENIOR":
        if mono_classe:
            c = CUSTO_SENIOR_MONO
        elif subord_zerado and (cand["qtd"] or cand["valor"]):
            c = CUSTO_SENIOR_SO_ELE
    if not cand["qtd"] and not cand["valor"]:
        c += CUSTO_ZERADA
    return c


def rotulo_classe(c):
    return f"{c['TIPO'] or '—'}/{c['SERIE'] or '—'}" + ("" if c["no"] == "SUBORD" else " [nó SENIOR]")


def escolher_serie_coerente(meses):
    """meses: {(ano, mes): dados de extrair_mes()}.
    Devolve (escolhas, diagnosticos, mono_classe), onde
      escolhas      = {(ano,mes): candidato escolhido}
      diagnosticos  = {(ano,mes): {"erro": float, "trocou": bool, "alt": str|None}}
        erro   = descontinuidade residual frente ao mês anterior escolhido
        trocou = a escolha difere da que o rótulo, isolado, indicaria
        alt    = rótulo que a regra antiga teria escolhido (quando trocou)
    """
    chaves = [k for k in sorted(meses) if meses[k].get("candidatos")]
    if not chaves:
        return {}, {}, False

    # Fundo "mono-classe": em nenhum mês os nós SENIOR e SUBORD têm valor ao
    # mesmo tempo. Nesses fundos há uma única classe de cotas e o nó em que o
    # administrador a colocou varia (ex.: SUPER RENDA — a mesma posição de
    # 10.176.737,63 cotas migrou de SENIOR para SUBORD em out/2024).
    def tem_valor(cs, no):
        return any(c["no"] == no and (c["qtd"] or c["valor"]) for c in cs)

    mono_classe = not any(tem_valor(meses[k]["candidatos"], "SENIOR")
                          and tem_valor(meses[k]["candidatos"], "SUBORD")
                          for k in chaves)

    # ---- Viterbi ----
    tabela = []          # (chave, candidatos, custo_acumulado, backpointer)
    custo_ant = cands_ant = chave_ant = None
    for k in chaves:
        cs = meses[k]["candidatos"]
        sub_zerado = not tem_valor(cs, "SUBORD")
        priors = [_custo_rotulo(c, sub_zerado, mono_classe) for c in cs]
        if custo_ant is None:
            custos, back = list(priors), [None] * len(cs)
        else:
            gap = (k[0] * 12 + k[1]) - (chave_ant[0] * 12 + chave_ant[1])
            custos, back = [], []
            for j, c in enumerate(cs):
                melhor = arg = None
                for p, cp in enumerate(cands_ant):
                    v = custo_ant[p] + PESO_CONTINUIDADE * min(
                        _erro_continuidade(cp, c, gap), ERRO_MAX)
                    if melhor is None or v < melhor:
                        melhor, arg = v, p
                custos.append(melhor + priors[j])
                back.append(arg)
        tabela.append((k, cs, custos, back))
        custo_ant, cands_ant, chave_ant = custos, cs, k

    # ---- reconstrução do caminho ----
    idx = min(range(len(tabela[-1][2])), key=lambda j: tabela[-1][2][j])
    caminho = [None] * len(tabela)
    for i in range(len(tabela) - 1, -1, -1):
        caminho[i] = idx
        if i:
            idx = tabela[i][3][idx]

    escolhas, diag = {}, {}
    for i, (k, cs, _c, _b) in enumerate(tabela):
        esc = cs[caminho[i]]
        escolhas[k] = esc
        # erro residual frente ao mês anterior escolhido
        if i:
            k_ant = tabela[i - 1][0]
            gap = (k[0] * 12 + k[1]) - (k_ant[0] * 12 + k_ant[1])
            erro = _erro_continuidade(escolhas[k_ant], esc, gap)
        else:
            erro = 0.0
        # o que a regra antiga (só rótulo) teria escolhido
        t_ant, s_ant = escolher_subordinada([c for c in cs if c["no"] == "SUBORD"])
        trocou = (t_ant, s_ant) != (esc["TIPO"], esc["SERIE"])
        diag[k] = {"erro": erro, "trocou": trocou,
                   "alt": (f"{t_ant or '—'}/{s_ant or '—'}" if trocou else None)}
    return escolhas, diag, mono_classe


def parse_informe(xml_bytes):
    """Compatibilidade: extrai um mês isolado, escolhendo a classe só pelo rótulo.
    O caminho usado pela interface é extrair_mes() + escolher_serie_coerente()."""
    d = extrair_mes(xml_bytes)
    subord = [c for c in d["candidatos"] if c["no"] == "SUBORD"]
    tipo_sel, serie_sel = escolher_subordinada(subord)
    d["sub_tipo"], d["sub_serie"] = tipo_sel, serie_sel
    esc = next((c for c in subord
                if (c["TIPO"], c["SERIE"]) == (tipo_sel, serie_sel)), None)
    d["sub_qtd"] = esc["qtd"] if esc else 0.0
    d["sub_valor"] = esc["valor"] if esc else 0.0
    d["aporte"] = esc["aporte"] if esc else 0.0
    d["resgate"] = esc["resgate"] if esc else 0.0
    return d


# ============================================================================
#  3. Leitura/escrita da planilha (cirurgia de XML — preserva gráficos/fórmulas)
# ============================================================================
def carregar_leitura(caminho, aba):
    return openpyxl.load_workbook(caminho, data_only=True)[aba]


def _add_meses(ano, mes, n):
    t = ano * 12 + (mes - 1) + n
    return t // 12, t % 12 + 1


def mapa_meses(caminho, aba, linha_data=18, linha_input=20):
    wv = openpyxl.load_workbook(caminho, data_only=True)[aba]
    mapa, vazio = {}, {}
    for c in range(4, wv.max_column + 1):
        v = wv.cell(linha_data, c).value
        if hasattr(v, "year"):
            k = (v.year, v.month)
            mapa[k] = get_column_letter(c)
            vazio[k] = wv.cell(linha_input, c).value is None
    if mapa:
        return mapa, vazio

    ancora = wv.cell(12, 8).value  # H12
    if not hasattr(ancora, "year"):
        raise RuntimeError(
            "Não há datas calculadas no cabeçalho (linha 18) e o 'Mês da Análise' (célula H12) "
            "está vazio. Preencha o Mês da Análise na planilha (ex.: 01/03/2026) — ou abra e "
            "salve o arquivo no Excel para recalcular — e rode novamente.")
    wf = openpyxl.load_workbook(caminho, data_only=False)[aba]
    cols = sorted(c for c in range(4, wf.max_column + 1)
                  if wf.cell(linha_data, c).value not in (None, ""))
    n = len(cols)
    for i, c in enumerate(cols):
        k = _add_meses(ancora.year, ancora.month, -(n - 1 - i))
        mapa[k] = get_column_letter(c)
        vazio[k] = wv.cell(linha_input, c).value is None
    return mapa, vazio


def celula_vazia(ws, coluna, linha):
    return ws[f"{coluna}{linha}"].value is None


def _col_idx(col):
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


def _fmt(v):
    f = float(v)
    if f == int(f) and abs(f) < 1e15:
        return str(int(f))
    return repr(f)


def _sheet_path(z, nome_aba):
    def attr(el, nome):
        m = re.search(fr'{nome}="([^"]+)"', el)
        return m.group(1) if m else None
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = {}
    for rel in re.findall(r"<Relationship\b[^>]*?>", rels):
        i, t = attr(rel, "Id"), attr(rel, "Target")
        if i and t:
            rid[i] = t
    for sh in re.findall(r"<sheet\b[^>]*?>", wb):
        if attr(sh, "name") == nome_aba:
            tgt = rid.get(attr(sh, "r:id"))
            if tgt:
                return "xl/" + tgt.lstrip("/") if not tgt.startswith("/") else tgt.lstrip("/")
    raise KeyError(f"aba '{nome_aba}' não encontrada no .xlsx")


def _editar_aba(xml, updates):
    byrow = defaultdict(dict)
    for ref, (kind, val) in updates.items():
        m = re.match(r"([A-Z]+)(\d+)", ref)
        byrow[m.group(2)][ref] = (m.group(1), kind, val)

    def repl(mrow):
        num, attrs, body = mrow.group("n"), mrow.group("attrs"), mrow.group("body")
        if num not in byrow:
            return mrow.group(0)
        cells = re.findall(r"<c\b[^>]*?(?:/>|>.*?</c>)", body, re.S)
        cmap = {re.search(r'r="([A-Z]+\d+)"', c).group(1): c for c in cells}
        for ref, (col, kind, val) in byrow[num].items():
            style = ""
            if ref in cmap:
                sm = re.search(r's="(\d+)"', cmap[ref])
                style = f' s="{sm.group(1)}"' if sm else ""
            else:
                melhor = None
                for r2, c2 in cmap.items():
                    if _col_idx(re.match(r"([A-Z]+)", r2).group(1)) < _col_idx(col):
                        sm = re.search(r's="(\d+)"', c2)
                        if sm:
                            melhor = sm.group(1)
                if melhor:
                    style = f' s="{melhor}"'
            if kind == "n":
                cmap[ref] = f'<c r="{ref}"{style}><v>{_fmt(val)}</v></c>'
            else:
                cmap[ref] = f'<c r="{ref}"{style}><f>{val}</f></c>'
        ordenado = sorted(cmap.values(),
                          key=lambda c: _col_idx(re.search(r'r="([A-Z]+)\d+"', c).group(1)))
        return f"<row{attrs}>" + "".join(ordenado) + "</row>"

    return re.sub(r'<row(?P<attrs>[^>]*\br="(?P<n>\d+)"[^>]*)>(?P<body>.*?)</row>',
                  repl, xml, flags=re.S)


def _forcar_recalc(wbxml):
    if "<calcPr" in wbxml:
        if "fullCalcOnLoad" in wbxml:
            return wbxml
        return re.sub(r"<calcPr\b([^>]*?)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wbxml)
    return re.sub(r"(</sheets>)", r'\1<calcPr calcId="0" fullCalcOnLoad="1"/>', wbxml, count=1)


def escrever_planilha(caminho, nome_aba, updates):
    caminho = str(caminho)
    with zipfile.ZipFile(caminho) as z:
        spath = _sheet_path(z, nome_aba)
        itens = {n: z.read(n) for n in z.namelist()}
    itens[spath] = _editar_aba(itens[spath].decode("utf-8"), updates).encode("utf-8")
    itens["xl/workbook.xml"] = _forcar_recalc(itens["xl/workbook.xml"].decode("utf-8")).encode("utf-8")
    tmp = caminho + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in itens.items():
            z.writestr(n, data)
    os.replace(tmp, caminho)


# ============================================================================
#  4. Orquestração da extração (mesma lógica do notebook), com log ao vivo
# ============================================================================
LINHAS_DADOS = {19: "sub_qtd", 20: "sub_valor", 22: "aporte", 23: "resgate",
                38: "pl", 43: "inadimplentes", 45: "inad_over90", 47: "recompras"}
MESES_PT = ["", "jan", "fev", "mar", "abr", "mai", "jun",
            "jul", "ago", "set", "out", "nov", "dez"]


def rotulo(k):
    return f"{MESES_PT[k[1]]}/{k[0]}"


def executar(cfg, log):
    """cfg: dict com modo, busca, excel, aba, meses, sobrescrever, backup.
    log(mensagem, nivel): callback de progresso (nivel: info/ok/erro/aviso/destaque)."""
    modo = cfg["modo"]
    busca = cfg["busca"]
    caminho = Path(cfg["excel"])
    aba = cfg["aba"]
    ultimos = int(cfg["meses"])
    sobrescrever = bool(cfg["sobrescrever"])
    backup = bool(cfg["backup"])

    if not busca.strip():
        raise RuntimeError("Informe o nome ou o CNPJ do fundo.")
    if not caminho.exists():
        raise RuntimeError(f"Planilha não encontrada:\n{caminho}")

    log(f"Modo de busca: {modo}", "info")
    log(f"Fundo (busca): {busca}", "info")
    log(f"Planilha: {caminho.name}", "info")
    log(f"Últimos meses: {'todos' if ultimos == 0 else ultimos}   "
        f"Sobrescrever: {'sim' if sobrescrever else 'não'}   "
        f"Backup: {'sim' if backup else 'não'}", "info")
    log("Conectando ao FNET…", "info")

    cli = FnetClient()
    cli.iniciar_sessao()
    id_fundo, nome_exato = cli.resolver_fundo(busca, modo=modo)
    log(f"Fundo resolvido: {nome_exato}  (idFundo={id_fundo})", "destaque")

    informes = cli.listar_informes(id_fundo, nome_exato)
    log(f"Informes mensais disponíveis no FNET: {len(informes)}", "info")

    ws = carregar_leitura(caminho, aba)
    mapa, vazio = mapa_meses(caminho, aba)
    log(f"Colunas de mês na planilha: {len(mapa)} "
        f"({rotulo(min(mapa))} … {rotulo(max(mapa))})", "info")

    resumo = {"preenchidos": [], "sem_informe": [], "ja_preenchidos": [], "falhas": [],
              "corrigidos": [], "saltos": []}
    updates = {}

    meses = sorted(mapa)
    if ultimos:
        meses = meses[-ultimos:]
    log(f"Janela considerada: {len(meses)} mês(es) "
        f"({rotulo(meses[0])} … {rotulo(meses[-1])})", "info")

    # ---- 1) baixa TODOS os meses da janela ----------------------------------
    # A janela inteira é necessária mesmo quando só alguns meses serão
    # gravados: a identificação da classe subordinada é feita comparando os
    # meses entre si (ver escolher_serie_coerente), então tirar meses do meio
    # enfraquece a análise.
    a_baixar = [k for k in meses if k in informes]
    for k in meses:
        if k not in informes:
            resumo["sem_informe"].append(k)
    log(f"Baixando {len(a_baixar)} informe(s) para análise de coerência…", "info")

    dados = {}
    for i, chave in enumerate(a_baixar, 1):
        try:
            dados[chave] = extrair_mes(cli.baixar_xml(informes[chave]))
        except Exception as e:
            resumo["falhas"].append((chave, str(e)))
            log(f"  FALHA ao baixar/ler {rotulo(chave)}: {e}", "erro")
        if i % 6 == 0 or i == len(a_baixar):
            log(f"  … {i}/{len(a_baixar)}", "info")

    sem_classe = [k for k, d in dados.items() if not d.get("candidatos")]
    for k in sem_classe:
        resumo["falhas"].append(
            (k, "o informe não traz a descrição das classes (formato antigo do FNET)"))
        log(f"  FALHA {rotulo(k)}: informe sem descrição de classes "
            "(formato antigo) — nada será gravado neste mês.", "erro")
        dados.pop(k)

    if not dados:
        raise RuntimeError("nenhum informe legível na janela escolhida.")

    # ---- 2) identifica a subordinada de forma coerente ----------------------
    log("Identificando a classe subordinada pela continuidade da série "
        "(quantidade de cotas, valor da cota, aportes e resgates)…", "info")
    escolhas, diag, mono = escolher_serie_coerente(dados)
    if mono:
        log("  Fundo com classe única: o informe alterna a posição entre os nós "
            "Sênior e Subordinada — a série foi reconstruída pela continuidade.", "aviso")

    # ---- 3) grava --------------------------------------------------------
    for chave in meses:
        if chave not in escolhas:
            continue
        col = mapa[chave]
        if not sobrescrever and not vazio[chave]:
            resumo["ja_preenchidos"].append(chave)
            continue
        d, c, dg = dados[chave], escolhas[chave], diag[chave]
        campos = dict(d, sub_qtd=c["qtd"], sub_valor=c["valor"],
                      aporte=c["aporte"], resgate=c["resgate"])
        for linha, campo in LINHAS_DADOS.items():
            updates[f"{col}{linha}"] = ("n", campos[campo])
        if sobrescrever or celula_vazia(ws, col, 41):
            updates[f"{col}41"] = ("f", f"{col}38-{col}43")
        resumo["preenchidos"].append(chave)

        extra = ""
        if dg["trocou"]:
            resumo["corrigidos"].append((chave, rotulo_classe(c), dg["alt"]))
            extra += f"  [corrigido: pelo rótulo seria {dg['alt']}]"
        if dg["erro"] > ERRO_ALERTA:
            resumo["saltos"].append((chave, dg["erro"]))
            extra += f"  [SALTO sem justificativa de aporte/resgate: {dg['erro']:.0%}]"
        log(f"  OK  {rotulo(chave):>8}  (col {col})  PL={d['pl']:,.2f}  "
            f"cotas={c['qtd']:,.4f}  valor={c['valor']:,.4f}  "
            f"sub={rotulo_classe(c)}{extra}",
            "aviso" if dg["erro"] > ERRO_ALERTA else "ok")

    if resumo["corrigidos"]:
        log(f"{len(resumo['corrigidos'])} mês(es) em que a classe do rótulo NÃO era a "
            "subordinada real (mezanino/sênior rotulado como subordinada) — "
            "corrigido pela continuidade da série.", "destaque")
    if resumo["saltos"]:
        log(f"{len(resumo['saltos'])} mês(es) com variação brusca que nem aporte nem "
            "resgate explicam — confira esses meses no FNET.", "aviso")

    if updates:
        if backup:
            bak = caminho.with_name(
                f"{caminho.stem}_backup_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
            shutil.copy2(caminho, bak)
            log(f"Backup criado: {bak.name}", "destaque")
        escrever_planilha(caminho, aba, updates)
        log(f"Planilha atualizada: {len(resumo['preenchidos'])} mês(es) preenchido(s).", "destaque")
    else:
        log("Nada a escrever (todos os meses já preenchidos ou sem informe).", "aviso")

    # resumo final estruturado
    resumo_txt = {
        "fundo": nome_exato,
        "idFundo": id_fundo,
        "planilha": str(caminho),
        "preenchidos": [rotulo(k) for k in sorted(resumo["preenchidos"])],
        "ja_preenchidos": [rotulo(k) for k in sorted(resumo["ja_preenchidos"])],
        "sem_informe": [rotulo(k) for k in sorted(resumo["sem_informe"])],
        "falhas": [f"{rotulo(k)}: {msg}" for k, msg in resumo["falhas"]],
        "corrigidos": [f"{rotulo(k)}: usou {novo} (o rótulo indicaria {velho})"
                       for k, novo, velho in sorted(resumo["corrigidos"])],
        "saltos": [f"{rotulo(k)}: variação de {e:.0%} sem aporte/resgate que a explique"
                   for k, e in sorted(resumo["saltos"])],
    }
    return resumo_txt


# ============================================================================
#  5. Interface HTML
# ============================================================================
PAGINA_HTML = r"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Extração FIDC · FNET — Netz Asset</title>
<style>
  :root{
    --navy:#001B5C; --navy2:#002F6C; --orange:#FF965A; --orange-d:#f57f3a;
    --blue:#116DFF; --bg:#eef3fa; --card:#ffffff; --line:#dfe6f1;
    --txt:#12233f; --muted:#5b6b8c; --ink:#001B5C;
    --con-bg:#04123a; --con-line:#12245e;
    --font:"IBM Plex Sans","Segoe UI",system-ui,Arial,sans-serif;
  }
  *{box-sizing:border-box}
  body{margin:0;font-family:var(--font);background:var(--bg);color:var(--txt)}
  .topbar{background:var(--navy)}
  .topbar .in{max-width:1000px;margin:0 auto;padding:16px 20px;display:flex;align-items:center;gap:16px}
  .brand{height:30px;width:auto;display:block}
  .topbar .tag{color:#aebfe0;font-size:13px;border-left:1px solid #24407e;padding-left:16px}
  .topbar .tag b{color:#fff}
  .wrap{max-width:1000px;margin:0 auto;padding:26px 20px 60px}
  .pagettl h1{font-size:20px;margin:0;font-weight:700;color:var(--ink)}
  .pagettl p{color:var(--muted);font-size:13px;margin:4px 0 0}
  .grid{display:grid;grid-template-columns:1fr 1fr;gap:16px;margin-top:22px}
  .card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:20px;box-shadow:0 1px 3px rgba(0,27,92,.06)}
  .card h2{font-size:12px;margin:0 0 16px;color:var(--navy);text-transform:uppercase;letter-spacing:.7px;font-weight:700}
  label{display:block;font-size:13px;margin:0 0 6px;color:var(--txt);font-weight:600}
  .hint{color:var(--muted);font-size:12px;margin:4px 0 0;font-weight:400}
  input[type=text],input[type=number],select{
    width:100%;padding:11px 12px;border-radius:9px;border:1px solid var(--line);
    background:#fff;color:var(--txt);font-size:14px;outline:none;font-family:var(--font)}
  input:focus,select:focus{border-color:var(--blue);box-shadow:0 0 0 3px rgba(17,109,255,.12)}
  .field{margin-bottom:16px}
  .seg{display:flex;background:#eaf0fa;border:1px solid var(--line);border-radius:9px;overflow:hidden;padding:3px;gap:3px}
  .seg button{flex:1;padding:9px;background:transparent;border:0;border-radius:7px;color:var(--navy);cursor:pointer;font-size:14px;font-weight:600;font-family:var(--font)}
  .seg button.active{background:var(--navy);color:#fff}
  .filerow{display:flex;gap:8px}
  .filerow input{flex:1}
  .btn{padding:11px 16px;border-radius:9px;border:1px solid var(--navy);background:#fff;
       color:var(--navy);cursor:pointer;font-size:14px;font-weight:600;font-family:var(--font);white-space:nowrap}
  .btn:hover{background:var(--navy);color:#fff}
  .switch{display:flex;align-items:center;justify-content:space-between;gap:14px;padding:12px 0;border-top:1px solid var(--line)}
  .switch:first-of-type{border-top:0}
  .switch > div > div{font-weight:600;color:var(--txt)}
  .toggle{position:relative;width:46px;height:26px;flex:0 0 auto}
  .toggle input{display:none}
  .track{position:absolute;inset:0;background:#c9d4e8;border-radius:20px;transition:.15s}
  .thumb{position:absolute;top:3px;left:3px;width:20px;height:20px;border-radius:50%;background:#fff;transition:.15s;box-shadow:0 1px 2px rgba(0,0,0,.25)}
  .toggle input:checked + .track{background:var(--navy)}
  .toggle input:checked + .track .thumb{left:23px}
  .go{width:100%;margin-top:10px;padding:14px;border-radius:11px;border:0;font-size:15px;font-weight:700;
      background:var(--orange);color:#3a1a00;cursor:pointer;font-family:var(--font)}
  .go:hover{background:var(--orange-d)}
  .go:disabled{opacity:.55;cursor:not-allowed}
  .console{margin-top:22px;background:var(--con-bg);border:1px solid var(--con-line);border-radius:14px;overflow:hidden}
  .console .bar{display:flex;align-items:center;gap:8px;padding:10px 14px;border-bottom:1px solid var(--con-line);color:#aebfe0;font-size:12px}
  .dot{width:10px;height:10px;border-radius:50%;background:#3a4a7a}
  .dot.on{background:#2fd4a7;box-shadow:0 0 8px #2fd4a7}
  .dot.busy{background:var(--orange);box-shadow:0 0 8px var(--orange)}
  .dot.err{background:#ff6b6b;box-shadow:0 0 8px #ff6b6b}
  pre#log{margin:0;padding:16px;height:340px;overflow:auto;font-family:"Cascadia Code",Consolas,monospace;font-size:13px;line-height:1.55;white-space:pre-wrap;color:#dbe5f5}
  .l-info{color:#cdd8ef} .l-ok{color:#48e0aa} .l-erro{color:#ff7b7b}
  .l-aviso{color:#ffc65c} .l-destaque{color:#ffb07d;font-weight:700}
  .foot{color:var(--muted);font-size:12px;margin-top:18px;text-align:center}
  @media(max-width:820px){.grid{grid-template-columns:1fr}}
</style>
</head>
<body>
<div class="topbar"><div class="in">
  <svg class="brand" viewBox="389.66 389.65 1407.66 300.69" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Netz Asset">
    <path fill="#ff965a" d="m537.274 537.267 73.8-73.8 73.801 73.8-73.8 73.8z"/>
    <path fill="#ffffff" d="m684.87 684.87 147.61-147.6V389.66L684.87 537.27z"/>
    <path fill="#ffffff" d="m389.66 684.87 147.61-147.6V389.66L389.66 537.27z"/>
    <path fill="#ffffff" d="M1096.23 437.44c-40.75 0-71.65 33.25-75.87 79.14v-73.52h-40.28v241.66h40.28V526.89c0-37.46 23.41-56.67 57.14-56.67 40.27 0 57.14 26.23 57.14 56.67v157.83h40.28V526.89c0-51.05-33.72-89.45-78.68-89.45Z"/>
    <path fill="#ffffff" d="M1319.63 437.43c-66.51 0-110.06 51.05-110.06 126.45s43.55 126.46 110.06 126.46c53.86 0 92.26-34.19 105.38-86.18h-41.69c-10.77 31.84-34.19 51.52-63.69 51.52-38.41 0-66.98-34.19-69.78-84.31h179.84c0-85.23-42.16-133.94-110.06-133.94m-67.44 101.17c7.49-40.28 33.72-66.51 67.44-66.51s59.95 26.23 67.44 66.51z"/>
    <path fill="#ffffff" d="M1527.56 389.66h-40.28v53.39h-41.21v34.66h41.21v147.52c0 38.88 22.02 59.48 63.7 59.48h17.8v-34.66h-18.27c-14.99 0-22.95-8.44-22.95-24.83V477.7h41.22v-34.66h-41.22v-53.39Z"/>
    <path fill="#ffffff" d="M1797.32 443.21h-193.89v34.5h176.85L1603.43 618.6v66.27h193.89v-34.81h-176.84L1797.32 514z"/>
  </svg>
  <span class="tag">Ferramenta interna &middot; <b>Extração FIDC</b></span>
</div></div>
<div class="wrap">
  <div class="pagettl">
    <h1>Extração FIDC · FNET</h1>
    <p>Preenche a aba <b>ANÁLISE FIDC</b> da planilha a partir dos Informes Mensais do FNET.</p>
  </div>

  <div class="grid">
    <!-- Coluna esquerda: identificação do fundo -->
    <div class="card">
      <h2>Fundo</h2>
      <div class="field">
        <label>Como identificar o fundo</label>
        <div class="seg" id="seg">
          <button type="button" data-modo="nome" class="active">Nome</button>
          <button type="button" data-modo="cnpj">CNPJ</button>
        </div>
        <p class="hint" id="modoHint">Digite o nome do fundo (o nome pode ser ambíguo).</p>
      </div>
      <div class="field">
        <label id="lblBusca">Nome do fundo</label>
        <input type="text" id="busca" placeholder="MULTIPLICA FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS">
      </div>
      <div class="field">
        <label>Planilha (arquivo .xlsx)</label>
        <div class="filerow">
          <input type="text" id="excel" placeholder="C:\...\Modelo.xlsx">
          <button type="button" class="btn" id="procurar">Procurar…</button>
        </div>
        <p class="hint">Clique em Procurar para escolher o arquivo no computador.</p>
      </div>
      <div class="field">
        <label>Aba da planilha</label>
        <input type="text" id="aba" value="ANÁLISE FIDC">
      </div>
    </div>

    <!-- Coluna direita: opções -->
    <div class="card">
      <h2>Opções</h2>
      <div class="field">
        <label>Últimos meses a considerar</label>
        <input type="number" id="meses" value="36" min="0">
        <p class="hint">Quantidade de colunas de mês (mais recentes). Use <b>0</b> para todas.</p>
      </div>
      <div class="switch">
        <div>
          <div>Sobrescrever meses já preenchidos</div>
          <p class="hint" style="margin:2px 0 0">Desligado: preenche só o que está em branco (recomendado).</p>
        </div>
        <label class="toggle"><input type="checkbox" id="sobrescrever"><span class="track"><span class="thumb"></span></span></label>
      </div>
      <div class="switch">
        <div>
          <div>Fazer backup da planilha</div>
          <p class="hint" style="margin:2px 0 0">Cria uma cópia com data/hora antes de gravar.</p>
        </div>
        <label class="toggle"><input type="checkbox" id="backup" checked><span class="track"><span class="thumb"></span></span></label>
      </div>
      <button class="go" id="go">▶  Extrair e preencher</button>
    </div>
  </div>

  <div class="console">
    <div class="bar"><span class="dot" id="dot"></span><span id="status">Pronto.</span></div>
    <pre id="log"></pre>
  </div>

  <p class="foot">Netz Asset · roda 100% no seu computador · dados vêm da API pública do FNET · o Excel é gravado localmente.</p>
</div>

<script>
  let modo = "nome";
  const $ = id => document.getElementById(id);
  const seg = $("seg"), logEl = $("log"), dot = $("dot"), status = $("status"), go = $("go");

  seg.addEventListener("click", e => {
    const b = e.target.closest("button"); if(!b) return;
    modo = b.dataset.modo;
    [...seg.children].forEach(x => x.classList.toggle("active", x === b));
    if(modo === "cnpj"){
      $("lblBusca").textContent = "CNPJ do fundo";
      $("busca").placeholder = "23.216.398/0001-01";
      $("modoHint").textContent = "Busca mais precisa (retorna 1 fundo). Pode digitar com ou sem pontuação.";
    } else {
      $("lblBusca").textContent = "Nome do fundo";
      $("busca").placeholder = "MULTIPLICA FUNDO DE INVESTIMENTO EM DIREITOS CREDITÓRIOS";
      $("modoHint").textContent = "Digite o nome do fundo (o nome pode ser ambíguo).";
    }
  });

  $("procurar").addEventListener("click", async () => {
    status.textContent = "Abrindo seletor de arquivo…";
    try{
      const r = await fetch("/browse");
      const j = await r.json();
      if(j.path) $("excel").value = j.path;
      status.textContent = j.path ? "Arquivo selecionado." : "Nenhum arquivo escolhido.";
    }catch(err){ status.textContent = "Não foi possível abrir o seletor: " + err; }
  });

  function addLine(text, nivel){
    const span = document.createElement("span");
    span.className = "l-" + (nivel || "info");
    span.textContent = text + "\n";
    logEl.appendChild(span);
    logEl.scrollTop = logEl.scrollHeight;
  }

  go.addEventListener("click", () => {
    const cfg = {
      modo, busca: $("busca").value, excel: $("excel").value, aba: $("aba").value,
      meses: $("meses").value || "0",
      sobrescrever: $("sobrescrever").checked ? "1" : "0",
      backup: $("backup").checked ? "1" : "0",
    };
    if(!cfg.busca.trim()){ addLine("⚠️  Informe o nome ou CNPJ do fundo.", "aviso"); return; }
    if(!cfg.excel.trim()){ addLine("⚠️  Selecione a planilha (.xlsx).", "aviso"); return; }

    logEl.innerHTML = "";
    go.disabled = true; dot.className = "dot busy"; status.textContent = "Processando…";
    const qs = new URLSearchParams(cfg).toString();
    const es = new EventSource("/run?" + qs);

    es.onmessage = ev => {
      const m = JSON.parse(ev.data);
      addLine(m.msg, m.nivel);
    };
    es.addEventListener("fim", ev => {
      es.close(); go.disabled = false; dot.className = "dot on"; status.textContent = "Concluído.";
      const r = JSON.parse(ev.data);
      addLine("", "info");
      addLine("──────── RESUMO ────────", "destaque");
      addLine("Fundo: " + r.fundo, "info");
      addLine("Preenchidos agora (" + r.preenchidos.length + "): " + (r.preenchidos.join(", ") || "—"), "ok");
      if(r.ja_preenchidos.length) addLine("Já preenchidos (" + r.ja_preenchidos.length + "): " + r.ja_preenchidos.join(", "), "info");
      if(r.sem_informe.length) addLine("Sem informe no FNET (" + r.sem_informe.length + "): " + r.sem_informe.join(", "), "aviso");
      if(r.corrigidos && r.corrigidos.length){
        addLine("", "info");
        addLine("Classe corrigida pela continuidade da série (" + r.corrigidos.length + "):", "destaque");
        r.corrigidos.forEach(f => addLine("  " + f, "destaque"));
      }
      if(r.saltos && r.saltos.length){
        addLine("", "info");
        addLine("Conferir manualmente — variação brusca sem justificativa (" + r.saltos.length + "):", "aviso");
        r.saltos.forEach(f => addLine("  " + f, "aviso"));
      }
      if(r.falhas.length){ addLine("Falhas (" + r.falhas.length + "):", "erro"); r.falhas.forEach(f => addLine("  " + f, "erro")); }
      addLine("", "info");
      addLine("Abra a planilha no Excel — as fórmulas recalculam automaticamente ao abrir.", "destaque");
    });
    es.addEventListener("erro", ev => {
      es.close(); go.disabled = false; dot.className = "dot err"; status.textContent = "Erro.";
      const e = JSON.parse(ev.data);
      addLine("❌  " + e.msg, "erro");
    });
    es.onerror = () => {
      if(go.disabled){ es.close(); go.disabled = false; dot.className = "dot err";
        status.textContent = "Conexão interrompida."; addLine("❌  Conexão com o servidor perdida.", "erro"); }
    };
  });
</script>
</body>
</html>
"""


# ============================================================================
#  6. Servidor HTTP local
# ============================================================================
def _sse(handler, evento, payload):
    """Envia um evento SSE. evento=None -> evento 'message' padrão."""
    linha = ""
    if evento:
        linha += f"event: {evento}\n"
    linha += "data: " + json.dumps(payload, ensure_ascii=False) + "\n\n"
    handler.wfile.write(linha.encode("utf-8"))
    handler.wfile.flush()


def escolher_arquivo():
    """Abre um seletor de arquivo nativo (tkinter). Retorna o caminho ou ''."""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        caminho = filedialog.askopenfilename(
            title="Selecione a planilha de análise",
            filetypes=[("Planilhas Excel", "*.xlsx"), ("Todos os arquivos", "*.*")])
        root.destroy()
        return caminho or ""
    except Exception as e:
        print("Seletor de arquivo indisponível:", e)
        return ""


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass  # silencia o log padrão

    def _html(self, body, code=200, ctype="text/html; charset=utf-8"):
        data = body.encode("utf-8") if isinstance(body, str) else body
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        rota = parsed.path

        if rota in ("/", "/index.html"):
            self._html(PAGINA_HTML)
            return

        if rota == "/browse":
            caminho = escolher_arquivo()
            self._html(json.dumps({"path": caminho}), ctype="application/json; charset=utf-8")
            return

        if rota == "/run":
            q = urllib.parse.parse_qs(parsed.query)
            cfg = {
                "modo": (q.get("modo", ["nome"])[0]),
                "busca": (q.get("busca", [""])[0]),
                "excel": (q.get("excel", [""])[0]),
                "aba": (q.get("aba", ["ANÁLISE FIDC"])[0]),
                "meses": (q.get("meses", ["0"])[0]),
                "sobrescrever": (q.get("sobrescrever", ["0"])[0] == "1"),
                "backup": (q.get("backup", ["1"])[0] == "1"),
            }
            self.send_response(200)
            self.send_header("Content-Type", "text/event-stream; charset=utf-8")
            self.send_header("Cache-Control", "no-cache")
            self.send_header("Connection", "keep-alive")
            self.end_headers()

            def log(msg, nivel="info"):
                _sse(self, None, {"msg": msg, "nivel": nivel})

            try:
                resumo = executar(cfg, log)
                _sse(self, "fim", resumo)
            except Exception as e:
                _sse(self, "erro", {"msg": str(e)})
            return

        self._html("404", code=404, ctype="text/plain; charset=utf-8")


def achar_porta(inicio=8765, fim=8815):
    import socket
    for p in range(inicio, fim):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("127.0.0.1", p)) != 0:
                return p
    return inicio


def main():
    porta = achar_porta()
    endereco = f"http://127.0.0.1:{porta}"
    servidor = ThreadingHTTPServer(("127.0.0.1", porta), Handler)
    print("=" * 60)
    print("  Extração FIDC · FNET — Netz Asset")
    print("=" * 60)
    print(f"  Servidor rodando em: {endereco}")
    print("  (O navegador deve abrir sozinho. Se não abrir, cole o")
    print("   endereço acima no navegador.)")
    print("  Para encerrar: feche esta janela ou pressione Ctrl+C.")
    print("=" * 60)
    if not os.environ.get("NETZ_HUB"):
        threading.Timer(0.8, lambda: webbrowser.open(endereco)).start()
    try:
        servidor.serve_forever()
    except KeyboardInterrupt:
        print("\nEncerrando…")
        servidor.shutdown()


if __name__ == "__main__":
    main()
